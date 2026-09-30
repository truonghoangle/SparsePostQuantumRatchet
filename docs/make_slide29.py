from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def tb(slide, left, top, width, height, text, size=14, bold=False, color=RGBColor(0xFF,0xFF,0xFF), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return tf

def ap(tf, text, size=14, bold=False, color=RGBColor(0xFF,0xFF,0xFF)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"

BG = RGBColor(0xFF, 0xFF, 0xFF)
WHITE = RGBColor(0x1E, 0x1E, 0x2E)
ACCENT = RGBColor(0x1A, 0x56, 0xDB)
HIGH = RGBColor(0xC0, 0x2B, 0x2B)
MED = RGBColor(0xC4, 0x6A, 0x10)
GREEN = RGBColor(0x1A, 0x7D, 0x3A)
DIM = RGBColor(0x55, 0x5B, 0x6E)

slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG

W = prs.slide_width
m = Inches(0.7)
col = Inches(5.6)

# Title
tf = tb(slide, m, Inches(0.4), W - 2*m, Inches(0.8),
    "Two Bugs Found by Formal Verification", 32, True, ACCENT)
ap(tf, "SPQR Lean Spec Audit", 16, False, DIM)

# Divider
line = slide.shapes.add_connector(1, m, Inches(1.3), W - m, Inches(1.3))
line.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
line.line.width = Pt(1)

# ---- LEFT: #1 HIGH ----
tf1 = tb(slide, m, Inches(1.6), col, Inches(4.8),
    "#1  Crash from Untrusted State Blob", 22, True, HIGH)
ap(tf1, "HIGH", 14, True, HIGH)
ap(tf1, "", 8)
ap(tf1, "The Problem", 15, True, ACCENT)
ap(tf1, "When the app reloads saved state, it trusts the data blindly.", 14, False, WHITE)
ap(tf1, "If the blob is corrupted or tampered with, the app crashes.", 14, False, WHITE)
ap(tf1, "", 10)
ap(tf1, "How the Spec Found It", 15, True, ACCENT)
ap(tf1, "We wrote specs for the loader and its consumers separately.", 14, False, DIM)
ap(tf1, "The loader promises nothing. The consumers require guarantees.", 14, False, DIM)
ap(tf1, "The gap between them is the bug \u2014 no code connects the two.", 14, False, DIM)
ap(tf1, "", 10)
ap(tf1, "Fix:  Check the data when loading. Reject bad blobs.", 14, True, GREEN)

# ---- RIGHT: #16 MEDIUM ----
rx = Inches(7.0)
tf2 = tb(slide, rx, Inches(1.6), col, Inches(4.8),
    "#16  Save/Load Loses Zero Polynomials", 22, True, MED)
ap(tf2, "MEDIUM", 14, True, MED)
ap(tf2, "", 8)
ap(tf2, "The Problem", 15, True, ACCENT)
ap(tf2, "Saving a zero polynomial produces empty bytes.", 14, False, WHITE)
ap(tf2, "Loading rejects empty bytes as invalid.", 14, False, WHITE)
ap(tf2, "So save \u2192 load is broken for any encoder with an empty slot.", 14, False, WHITE)
ap(tf2, "", 10)
ap(tf2, "How the Spec Found It", 15, True, ACCENT)
ap(tf2, "Three specs chain together: \u201czero \u2192 empty\u201d, \u201cserialize \u2192 bytes\u201d,", 14, False, DIM)
ap(tf2, "\u201cdeserialize rejects empty\u201d. The contradiction is immediate.", 14, False, DIM)
ap(tf2, "V1 never hits this (messages large enough), but the API is broken.", 13, False, DIM)
ap(tf2, "", 10)
ap(tf2, "Fix:  Accept empty bytes as the zero polynomial.", 14, True, GREEN)

# Note: not reachable from Signal
NOTE = RGBColor(0x6C, 0x6C, 0x6C)
tf_note = tb(slide, m, Inches(5.7), W - 2*m, Inches(0.8),
    "Not reachable from Signal:", 12, True, DIM)
ap(tf_note, "#1 — Signal\u2019s MAC-before-commit pattern authenticates every inbound message before state is committed. "
    "A remote attacker cannot inject a corrupted state blob through the network \u2014 it requires local access to the device\u2019s persisted storage.", 11, False, NOTE)
ap(tf_note, "#16 — The V1 protocol only encodes messages \u226596 bytes, which always fill all 16 polynomial slots. "
    "The zero-polynomial path is unreachable in Signal\u2019s usage; only direct API callers with very short messages can trigger it.", 11, False, NOTE)

# Bottom
tb(slide, m, Inches(6.8), W - 2*m, Inches(0.5),
    "Takeaway:  The math tells you exactly where the code breaks its own promises.",
    14, True, GREEN, PP_ALIGN.CENTER)

prs.save("/Users/hoangtruong/Documents/lean/Max/SparsePostQuantumRatchet/docs/report29.pptx")
print("Done")
